import re

from django.core.management.base import BaseCommand
from openpyxl import load_workbook

from gebarometer.models import GebarometerItem
from recordings.models import SignbankEntry


class Command(BaseCommand):
    help = "Importeer historische Gebarometer-data uit een Excelbestand."

    def add_arguments(self, parser):
        parser.add_argument(
            "file_path",
            type=str,
            help="Pad naar het Excelbestand.",
        )

    def handle(self, *args, **options):
        file_path = options["file_path"]

        workbook = load_workbook(
            filename=file_path,
            data_only=True,
            read_only=True,
        )

        worksheet = workbook.active

        headers = [
            cell.value
            for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
        ]

        header_map = {
            header: index
            for index, header in enumerate(headers)
            if header
        }

        required_headers = [
            "Datum",
            "Glossary",
            "Totaal Stemmen",
            "JA - JA",
            "JA - NEE",
            "NEE",
            "TWIJFEL",
            "Opmerking",
            "(Interne Opmerkingen)",
            "O-VL",
            "W-VL",
            "VL-B",
            "A",
            "L",
            "MailChimpLink",
        ]

        missing_headers = [
            header
            for header in required_headers
            if header not in header_map
        ]

        if missing_headers:
            self.stderr.write(
                self.style.ERROR(
                    "Ontbrekende kolommen: "
                    + ", ".join(missing_headers)
                )
            )
            return

        created_count = 0
        updated_count = 0
        skipped_count = 0
        linked_count = 0
        unlinked_count = 0

        def value(row, column):
            return row[header_map[column]]

        def integer(value):
            if value in (None, ""):
                return 0

            try:
                return int(value)
            except (TypeError, ValueError):
                return 0

        def text_value(value):
            if value is None:
                return ""

            return str(value).strip()

        def extract_report_id(url):
            if not url:
                return None

            match = re.search(r"[?&]id=([^&]+)", url)

            if not match:
                return None

            return match.group(1)

        for excel_row_number, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            planned_date = value(row, "Datum")
            gloss = text_value(value(row, "Glossary"))

            if not planned_date or not gloss or gloss.startswith("-"):
                skipped_count += 1
                continue

            if hasattr(planned_date, "date"):
                planned_date = planned_date.date()

            mailchimp_url = text_value(
                value(row, "MailChimpLink")
            )

            mailchimp_report_id = extract_report_id(
                mailchimp_url
            )

            signbank_entry = (
                SignbankEntry.objects
                .filter(gloss__iexact=gloss)
                .first()
            )

            if signbank_entry:
                linked_count += 1
            else:
                unlinked_count += 1

            defaults = {
                "signbank_entry": signbank_entry,
                "historical_gloss": gloss,
                "mailchimp_status": (
                    "VERSTUURD"
                    if mailchimp_url
                    else "NIET_VERSTUURD"
                ),
                "mailchimp_report_id": mailchimp_report_id,
                "mailchimp_url": mailchimp_url,
                "total_votes": integer(
                    value(row, "Totaal Stemmen")
                ),
                "ja_ja": integer(
                    value(row, "JA - JA")
                ),
                "ja_nee": integer(
                    value(row, "JA - NEE")
                ),
                "nee": integer(
                    value(row, "NEE")
                ),
                "twijfel": integer(
                    value(row, "TWIJFEL")
                ),
                "oost_vlaanderen": integer(
                    value(row, "O-VL")
                ),
                "west_vlaanderen": integer(
                    value(row, "W-VL")
                ),
                "vlaams_brabant": integer(
                    value(row, "VL-B")
                ),
                "antwerpen": integer(
                    value(row, "A")
                ),
                "limburg": integer(
                    value(row, "L")
                ),
                "remarks": text_value(
                    value(row, "Opmerking")
                ),
                "internal_remarks": text_value(
                    value(row, "(Interne Opmerkingen)")
                ),
            }

            try:
                item, created = (
                    GebarometerItem.objects.update_or_create(
                        planned_date=planned_date,
                        historical_gloss=gloss,
                        defaults=defaults,
                    )
                )
            except Exception as exc:
                self.stderr.write(
                    self.style.ERROR(
                        f"Rij {excel_row_number}: "
                        f"{gloss} kon niet worden geïmporteerd: "
                        f"{exc}"
                    )
                )
                skipped_count += 1
                continue

            if created:
                created_count += 1
            else:
                updated_count += 1

        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS(
                "Gebarometer-import voltooid."
            )
        )

        self.stdout.write(
            f"Nieuw aangemaakt: {created_count}"
        )
        self.stdout.write(
            f"Bijgewerkt: {updated_count}"
        )
        self.stdout.write(
            f"Overgeslagen: {skipped_count}"
        )
        self.stdout.write(
            f"Gekoppeld aan Signbank: {linked_count}"
        )
        self.stdout.write(
            f"Geen Signbank-match: {unlinked_count}"
        )